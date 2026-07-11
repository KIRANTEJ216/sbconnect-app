from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("output/presentation/SB_Connect_Client_Feature_Deck.pptx")
LOGO = Path("public/sbconnect-logo.png")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

W, H = prs.slide_width, prs.slide_height

COLORS = {
    "ink": RGBColor(24, 32, 52),
    "muted": RGBColor(95, 108, 130),
    "line": RGBColor(222, 228, 239),
    "canvas": RGBColor(248, 250, 253),
    "white": RGBColor(255, 255, 255),
    "primary": RGBColor(47, 63, 184),
    "primary_dark": RGBColor(34, 47, 145),
    "violet": RGBColor(181, 54, 197),
    "green": RGBColor(22, 134, 90),
    "gold": RGBColor(201, 145, 46),
    "red": RGBColor(196, 65, 63),
    "soft_blue": RGBColor(239, 243, 255),
    "soft_green": RGBColor(232, 248, 240),
    "soft_gold": RGBColor(255, 247, 230),
    "soft_red": RGBColor(255, 239, 238),
}


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def set_line(shape, color=COLORS["line"], width=1):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def no_line(shape):
    shape.line.fill.background()


def add_bg(slide, color=COLORS["canvas"]):
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, H)
    set_fill(bg, color)
    no_line(bg)
    bg.z_order = 0


def text_box(slide, text, x, y, w, h, size=18, color=COLORS["ink"], bold=False,
             align=PP_ALIGN.LEFT, font="Arial", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None, kicker=None):
    if kicker:
      text_box(slide, kicker.upper(), 0.72, 0.42, 4.8, 0.25, 8.5, COLORS["primary"], True)
    text_box(slide, title, 0.7, 0.72, 8.4, 0.55, 24, COLORS["ink"], True)
    if subtitle:
        text_box(slide, subtitle, 0.72, 1.27, 9.4, 0.36, 10.5, COLORS["muted"])
    add_footer(slide)


def add_footer(slide):
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.7), Inches(7.08), Inches(11.95), Inches(0.01))
    set_fill(line, COLORS["line"])
    no_line(line)
    text_box(slide, "SB Connect - Client Feature Deck", 0.72, 7.15, 3.2, 0.2, 7.5, COLORS["muted"])
    text_box(slide, "No Politics Only Business", 10.0, 7.15, 2.6, 0.2, 7.5, COLORS["muted"], align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, body, accent=COLORS["primary"], fill=COLORS["white"], icon=None):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shp, fill)
    set_line(shp, COLORS["line"], 1)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.08))
    set_fill(bar, accent)
    no_line(bar)
    if icon:
        circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.16), Inches(y + 0.18), Inches(0.42), Inches(0.42))
        set_fill(circ, accent)
        no_line(circ)
        text_box(slide, icon, x + 0.255, y + 0.255, 0.22, 0.14, 11, COLORS["white"], True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        title_x = x + 0.68
        title_w = w - 0.82
    else:
        title_x = x + 0.2
        title_w = w - 0.4
    text_box(slide, title, title_x, y + 0.2, title_w, 0.34, 13, COLORS["ink"], True)
    text_box(slide, body, x + 0.22, y + 0.68, w - 0.44, h - 0.86, 9.2, COLORS["muted"])
    return shp


def pill(slide, text, x, y, w, color=COLORS["soft_blue"], text_color=COLORS["primary"]):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.32))
    set_fill(shp, color)
    set_line(shp, color, 1)
    text_box(slide, text, x + 0.08, y + 0.075, w - 0.16, 0.12, 8.5, text_color, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_logo(slide, x=0.68, y=0.42, h=0.55, tile=False):
    if LOGO.exists():
        if tile:
            pad = 0.12
            tile_shape = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                Inches(x - pad),
                Inches(y - pad),
                Inches(h + 0.24),
                Inches(h + 0.24),
            )
            set_fill(tile_shape, COLORS["white"])
            no_line(tile_shape)
        slide.shapes.add_picture(str(LOGO), Inches(x), Inches(y), height=Inches(h))


def workflow_step(slide, n, title, body, x, y, w=2.15, accent=COLORS["primary"]):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.04))
    set_fill(shp, COLORS["white"])
    set_line(shp, COLORS["line"], 1)
    circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.16), Inches(y + 0.18), Inches(0.34), Inches(0.34))
    set_fill(circ, accent)
    no_line(circ)
    text_box(slide, str(n), x + 0.255, y + 0.24, 0.15, 0.12, 8, COLORS["white"], True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, title, x + 0.6, y + 0.16, w - 0.76, 0.24, 10.5, COLORS["ink"], True)
    text_box(slide, body, x + 0.18, y + 0.55, w - 0.36, 0.3, 8.0, COLORS["muted"])


def arrow(slide, x, y, w=0.38):
    text_box(slide, "->", x, y, w, 0.25, 13, COLORS["muted"], True, align=PP_ALIGN.CENTER)


def simple_table(slide, x, y, w, h, headers, rows, col_fracs):
    row_h = h / (len(rows) + 1)
    cx = x
    for i, head in enumerate(headers):
        cw = w * col_fracs[i]
        shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(cx), Inches(y), Inches(cw), Inches(row_h))
        set_fill(shp, COLORS["soft_blue"])
        set_line(shp, COLORS["line"], 0.7)
        text_box(slide, head, cx + 0.12, y + 0.12, cw - 0.22, row_h - 0.16, 8.2, COLORS["ink"], True)
        cx += cw
    for r_i, row in enumerate(rows):
        cx = x
        for c_i, cell in enumerate(row):
            cw = w * col_fracs[c_i]
            shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(cx), Inches(y + row_h * (r_i + 1)), Inches(cw), Inches(row_h))
            set_fill(shp, COLORS["white"])
            set_line(shp, COLORS["line"], 0.6)
            text_box(slide, cell, cx + 0.12, y + row_h * (r_i + 1) + 0.11, cw - 0.22, row_h - 0.12, 7.3, COLORS["muted"])
            cx += cw


def add_phone_mock(slide, x, y, w, h):
    phone = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(phone, RGBColor(22, 27, 45))
    no_line(phone)
    screen = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x + 0.12), Inches(y + 0.18), Inches(w - 0.24), Inches(h - 0.36))
    set_fill(screen, COLORS["canvas"])
    no_line(screen)
    text_box(slide, "Dashboard", x + 0.32, y + 0.38, w - 0.64, 0.25, 11, COLORS["ink"], True)
    for i, (label, color) in enumerate([("Membership", COLORS["green"]), ("Requests", COLORS["gold"]), ("Notifications", COLORS["red"])]):
        yy = y + 0.84 + i * 0.62
        mini = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x + 0.32), Inches(yy), Inches(w - 0.64), Inches(0.44))
        set_fill(mini, COLORS["white"])
        set_line(mini, COLORS["line"], 0.6)
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.48), Inches(yy + 0.14), Inches(0.16), Inches(0.16))
        set_fill(dot, color)
        no_line(dot)
        text_box(slide, label, x + 0.75, yy + 0.14, w - 1.1, 0.1, 7.5, COLORS["muted"], True)
    text_box(slide, "Quick Actions", x + 0.32, y + 2.96, w - 0.64, 0.2, 8.3, COLORS["ink"], True)
    pill(slide, "Create Request", x + 0.32, y + 3.26, w - 0.64, COLORS["soft_blue"], COLORS["primary"])
    pill(slide, "Record Deal", x + 0.32, y + 3.68, w - 0.64, COLORS["soft_green"], COLORS["green"])


def new_slide(bg=COLORS["canvas"]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    return slide


# 1 Cover
slide = new_slide(COLORS["primary_dark"])
circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9.2), Inches(4.35), Inches(4.8), Inches(4.8))
set_fill(circle, COLORS["violet"])
no_line(circle)
circle2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(-1.0), Inches(-1.1), Inches(3.8), Inches(3.8))
set_fill(circle2, COLORS["primary"])
no_line(circle2)
add_logo(slide, 0.82, 0.66, 0.62, tile=True)
text_box(slide, "SB Connect", 0.82, 1.88, 6.5, 0.62, 35, COLORS["white"], True)
text_box(slide, "Client Feature Presentation", 0.86, 2.62, 5.7, 0.32, 16, RGBColor(230, 235, 255), True)
text_box(slide, "A business community portal for member profiles, opportunity requests, meetings, attendance, admin operations, and future payments.", 0.88, 3.22, 7.1, 0.8, 15, RGBColor(238, 241, 255))
pill(slide, "No Politics Only Business", 0.88, 4.32, 2.35, RGBColor(238, 241, 255), COLORS["primary"])
text_box(slide, "Prepared for client review - July 2026", 0.88, 6.78, 4.2, 0.28, 9.2, RGBColor(226, 232, 255))

# 2 Product at a glance
slide = new_slide()
add_title(slide, "Product at a glance", "SB Connect brings membership, business discovery, opportunities, attendance, and admin governance into one portal.", "Overview")
card(slide, 0.75, 1.85, 2.7, 1.38, "Member Portal", "Secure login, profile setup, directory browsing, requests, chats, RSVPs, attendance, and deal recording.", COLORS["primary"], icon="M")
card(slide, 3.75, 1.85, 2.7, 1.38, "Business Exchange", "Members post requirements, pitch for work, award contracts, and track value generated in the network.", COLORS["green"], icon="B")
card(slide, 6.75, 1.85, 2.7, 1.38, "Meeting Hub", "Meeting announcements, RSVP, QR attendance links, attendance history, and compliance visibility.", COLORS["gold"], icon="Q")
card(slide, 9.75, 1.85, 2.7, 1.38, "Admin Console", "Approvals, roles, meetings, notifications, issue reports, exports, audit reports, and health checks.", COLORS["violet"], icon="A")
simple_table(slide, 1.0, 4.0, 11.25, 1.55, ["Who uses it?", "What they do", "Outcome"], [
    ["Members", "Create profiles, find businesses, post needs, pitch, attend meetings", "More business visibility and participation"],
    ["Admins", "Verify members, manage meetings, review reports, export data", "Centralized community operations"],
    ["Leadership", "Review business value, attendance, activity, and roadmap", "Better governance and growth decisions"],
], [0.22, 0.49, 0.29])

# 3 User journey
slide = new_slide()
add_title(slide, "End-user journey", "A simple path from account creation to active business participation.", "Member experience")
steps = [
    ("Register", "Create account with name, email, phone, password."),
    ("Create Profile", "Add company, categories, keywords, photo, catalog."),
    ("Get Verified", "Admin verifies profile and member status."),
    ("Discover", "Search directory and browse opportunities."),
    ("Participate", "Pitch, chat, RSVP, attend, and record deals."),
]
x = 0.72
for i, (t, b) in enumerate(steps, 1):
    workflow_step(slide, i, t, b, x + (i - 1) * 2.46, 2.25, 2.1, COLORS["primary"] if i < 5 else COLORS["green"])
    if i < 5:
        arrow(slide, x + (i - 1) * 2.46 + 2.12, 2.62)
text_box(slide, "The app is designed as an operational member workspace, not a static brochure. Users land on the dashboard and immediately see membership status, requests, meetings, notifications, and quick actions.", 1.05, 4.55, 11.2, 0.8, 14, COLORS["ink"])

# 4 Dashboard
slide = new_slide()
add_title(slide, "Dashboard and daily member workspace", "The dashboard gives members the most important actions and community signals in one place.", "Core screen")
add_phone_mock(slide, 0.9, 1.55, 2.25, 4.85)
card(slide, 3.65, 1.72, 2.55, 1.12, "Membership status", "Shows active, expired, or inactive state plus renewal countdown.", COLORS["green"])
card(slide, 6.45, 1.72, 2.55, 1.12, "Requests snapshot", "Highlights own requests, open opportunities, and new request indicators.", COLORS["gold"])
card(slide, 9.25, 1.72, 2.55, 1.12, "Notifications", "Unread alerts for issue resolution and new pitches.", COLORS["red"])
card(slide, 3.65, 3.25, 2.55, 1.12, "Quick actions", "Create request, view profile, and record business given.", COLORS["primary"])
card(slide, 6.45, 3.25, 2.55, 1.12, "Leaderboard", "Ranks contributors by business value and deal count.", COLORS["violet"])
card(slide, 9.25, 3.25, 2.55, 1.12, "Meetings", "Upcoming meetings, RSVP buttons, and attendance warnings.", COLORS["green"])
text_box(slide, "Client value: members need fewer instructions because the dashboard continuously points them toward the next useful action.", 3.65, 5.25, 8.4, 0.5, 13.5, COLORS["ink"], True)

# 5 Profiles and directory
slide = new_slide()
add_title(slide, "Business profiles and member directory", "Each member gets a searchable profile that represents their business inside the community.", "Discovery")
card(slide, 0.8, 1.74, 2.7, 1.28, "Profile creation", "Owner details, phone, company name, categories, size, location, email, website, and description.", COLORS["primary"])
card(slide, 3.72, 1.74, 2.7, 1.28, "Rich media", "Profile photo plus up to five catalog files such as PDF or images.", COLORS["violet"])
card(slide, 6.64, 1.74, 2.7, 1.28, "Search signals", "Categories, location, keywords, membership status, and verification badge.", COLORS["green"])
card(slide, 9.56, 1.74, 2.7, 1.28, "QR sharing", "Each profile has a QR-ready profile URL for offline networking.", COLORS["gold"])
simple_table(slide, 0.9, 3.9, 11.65, 1.72, ["Directory search", "Member benefit", "Admin benefit"], [
    ["Company name, category, keyword, location", "Find credible vendors, partners, and service providers", "Cleaner network visibility and stronger business matching"],
    ["Verified and membership status badges", "Trust signals while browsing", "Governance is visible without extra explanation"],
], [0.32, 0.36, 0.32])

# 6 Requests and deals
slide = new_slide()
add_title(slide, "Requests, pitches, and business value", "The request marketplace turns member needs into trackable business opportunities.", "Opportunities")
workflow_step(slide, 1, "Post request", "Title, category, description, budget, deadline.", 0.78, 2.05, 2.05, COLORS["primary"])
arrow(slide, 2.85, 2.42)
workflow_step(slide, 2, "Browse", "Members filter by category and status.", 3.28, 2.05, 2.05, COLORS["gold"])
arrow(slide, 5.35, 2.42)
workflow_step(slide, 3, "Pitch", "Interested member sends message and contact.", 5.78, 2.05, 2.05, COLORS["green"])
arrow(slide, 7.85, 2.42)
workflow_step(slide, 4, "Award", "Owner/admin awards and closes request.", 8.28, 2.05, 2.05, COLORS["violet"])
arrow(slide, 10.35, 2.42)
workflow_step(slide, 5, "Track value", "Deal feeds leaderboard and reports.", 10.78, 2.05, 2.05, COLORS["green"])
card(slide, 1.0, 4.25, 3.5, 1.15, "For members", "More qualified opportunities, direct phone follow-up, and chat-based networking.", COLORS["primary"])
card(slide, 4.9, 4.25, 3.5, 1.15, "For organizers", "Measurable deal creation and proof of community impact.", COLORS["green"])
card(slide, 8.8, 4.25, 3.5, 1.15, "For reporting", "Deal count and business value are visible through leaderboard and audits.", COLORS["gold"])

# 7 Meetings and attendance
slide = new_slide()
add_title(slide, "Meetings, RSVP, and attendance", "SB Connect supports recurring community meetings with digital attendance workflows.", "Participation")
card(slide, 0.85, 1.8, 2.7, 1.25, "Meeting creation", "Admins create meeting date, label, location, active state, and QR attendance URL.", COLORS["primary"])
card(slide, 3.75, 1.8, 2.7, 1.25, "Member RSVP", "Members respond Yes or No from upcoming meeting cards on dashboard.", COLORS["green"])
card(slide, 6.65, 1.8, 2.7, 1.25, "Attendance marking", "Members mark attendance or scan QR link; duplicate attendance is prevented.", COLORS["gold"])
card(slide, 9.55, 1.8, 2.7, 1.25, "Compliance", "A six-month attendance rule checks whether a member attended at least three meetings.", COLORS["red"])
simple_table(slide, 1.0, 4.02, 11.25, 1.5, ["End user", "Admin", "Leadership"], [
    ["Knows upcoming meetings and attendance history", "Sees attendance and RSVP lists with CSV export", "Tracks engagement and member participation"],
    ["Receives warnings if compliance is low", "Gets QR code for meeting attendance", "Can align benefits with active participation"],
], [0.34, 0.34, 0.32])

# 8 Communication and support
slide = new_slide()
add_title(slide, "Communication, support, and notifications", "The app includes direct messaging and a built-in issue reporting loop.", "Support")
card(slide, 0.95, 1.7, 2.9, 1.32, "Member chat", "One-to-one realtime conversations, unread counts, read status, and 30-day cleanup.", COLORS["primary"])
card(slide, 4.05, 1.7, 2.9, 1.32, "User notifications", "New pitch alerts and issue resolution messages appear in the dashboard.", COLORS["red"])
card(slide, 7.15, 1.7, 2.9, 1.32, "Report issue", "Floating support button lets users submit page, subject, and issue description.", COLORS["gold"])
card(slide, 10.25, 1.7, 2.1, 1.32, "Resolution loop", "Admins resolve reports and notify users with notes.", COLORS["green"])
text_box(slide, "Why this matters to the client", 1.0, 4.15, 4.0, 0.3, 18, COLORS["ink"], True)
text_box(slide, "Support is built into the operating experience. Members do not need a separate helpdesk link to report problems, and admins can close the loop from the admin panel.", 1.0, 4.58, 10.8, 0.68, 15, COLORS["muted"])

# 9 Admin console
slide = new_slide()
add_title(slide, "Admin console and operations", "Admins have one place to manage members, meetings, requests, reports, exports, and system health.", "Administration")
items = [
    ("Approvals", "Verify business profiles and manage pending members.", COLORS["green"]),
    ("Requests", "Close, delete, or award deals from business requests.", COLORS["primary"]),
    ("Meetings", "Create meetings, view QR, RSVPs, attendance, and CSV exports.", COLORS["gold"]),
    ("Notifications", "Publish announcements and meeting messages.", COLORS["violet"]),
    ("Issue Reports", "Review, resolve, notify user, or delete submitted reports.", COLORS["red"]),
    ("Audit & Health", "Generate audit reports and run data/system diagnostics.", COLORS["green"]),
    ("Webhook Sync", "Export app data to webhook/Google Sheets automation.", COLORS["primary"]),
    ("Roles", "Promote, view, and manage admin users.", COLORS["violet"]),
]
for idx, (title, body, color) in enumerate(items):
    row, col = divmod(idx, 4)
    card(slide, 0.72 + col * 3.12, 1.7 + row * 1.65, 2.72, 1.18, title, body, color)

# 10 Controls and UX
slide = new_slide()
add_title(slide, "Implemented controls and user experience", "The interface uses familiar controls that support fast member workflows.", "UX inventory")
simple_table(slide, 0.85, 1.65, 11.65, 3.7, ["Control group", "Examples", "User value"], [
    ["Forms", "Login, register, profile, request, meeting, webhook, issue report", "Users can complete work without leaving the app"],
    ["Selection", "Category buttons, dropdowns, date inputs, filter pills", "Simple filtering and structured data entry"],
    ["Media", "Photo upload, catalog upload, QR profile and attendance links", "Profiles become richer and meetings work offline"],
    ["Feedback", "Badges, loading states, errors, success messages, notifications", "Members understand status and next steps"],
    ["Navigation", "Desktop sidebar, mobile drawer, bottom navigation, clickable cards", "Responsive use across desktop and mobile"],
    ["Admin actions", "CSV export, collapsible sections, confirmations, role tools", "Operational tasks are organized and protected"],
], [0.22, 0.43, 0.35])

# 11 Roadmap
slide = new_slide()
add_title(slide, "Recommended next steps", "The current app is a strong MVP; these additions would make it production-ready and easier to scale.", "Roadmap")
roadmap = [
    ("1", "Complete payments", "Finish Razorpay membership payments, renewals, receipts, and reconciliation.", COLORS["green"]),
    ("2", "Move sensitive actions server-side", "Use Cloud Functions for role changes, attendance validation, deal awarding, and payment webhooks.", COLORS["primary"]),
    ("3", "Add analytics", "Dashboards for member growth, attendance trends, RSVP conversion, request closure, and deal value.", COLORS["violet"]),
    ("4", "Strengthen data quality", "Normalize money values, server timestamps, backend cleanup jobs, and stronger validation rules.", COLORS["gold"]),
    ("5", "Expand testing", "Automated tests for auth, profile, requests, duplicate prevention, attendance, and admin permissions.", COLORS["red"]),
]
for idx, (num, title, body, color) in enumerate(roadmap):
    y = 1.55 + idx * 0.92
    circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.95), Inches(y + 0.04), Inches(0.38), Inches(0.38))
    set_fill(circ, color)
    no_line(circ)
    text_box(slide, num, 1.07, y + 0.115, 0.14, 0.12, 8, COLORS["white"], True, align=PP_ALIGN.CENTER)
    text_box(slide, title, 1.55, y, 3.25, 0.25, 13.5, COLORS["ink"], True)
    text_box(slide, body, 5.08, y + 0.01, 7.2, 0.35, 11.2, COLORS["muted"])

# 12 Client takeaways
slide = new_slide(COLORS["primary_dark"])
circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9.7), Inches(-0.75), Inches(4.6), Inches(4.6))
set_fill(circle, COLORS["violet"])
no_line(circle)
add_logo(slide, 0.82, 0.66, 0.58, tile=True)
text_box(slide, "Client takeaways", 0.85, 1.55, 5.8, 0.55, 30, COLORS["white"], True)
takeaways = [
    "SB Connect is ready to present as a practical membership and business networking MVP.",
    "The strongest user value is the combination of directory, requests, meetings, attendance, and deal tracking.",
    "Admin operations are broad enough for real community management, including support reports and data exports.",
    "The best next investment is production hardening: payments, backend validation, analytics, and tests.",
]
for i, t in enumerate(takeaways):
    y = 2.55 + i * 0.78
    circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(1.0), Inches(y + 0.05), Inches(0.22), Inches(0.22))
    set_fill(circ, RGBColor(238, 241, 255))
    no_line(circ)
    text_box(slide, t, 1.42, y, 9.4, 0.34, 15, RGBColor(238, 241, 255))
pill(slide, "Business Network - Membership - Attendance - Admin Operations", 0.98, 6.28, 4.9, RGBColor(238, 241, 255), COLORS["primary"])

prs.save(OUT)
print(OUT)
